"""Generate executive presentation deck for TAMY - Neo Medallion Thesis Monitor."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    c_navy = RGBColor(15, 23, 42)       # #0f172a
    c_blue = RGBColor(26, 86, 219)      # #1a56db
    c_light_bg = RGBColor(248, 250, 252) # #f8fafc
    c_card_bg = RGBColor(255, 255, 255)
    c_text_dark = RGBColor(30, 41, 59)  # #1e293b
    c_text_muted = RGBColor(100, 116, 139) # #64748b
    c_border = RGBColor(226, 232, 240)  # #e2e8f0
    c_accent_green = RGBColor(22, 163, 74)
    c_accent_amber = RGBColor(217, 119, 6)
    c_accent_red = RGBColor(220, 38, 38)

    def add_header(slide, title, category="TAMY · HACKATHON SMC 2026"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_blue

        # Header title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = c_navy

    # -------------------------------------------------------------
    # SLIDE 1: Capa
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_navy
    bg1.line.color.rgb = c_navy

    # Tagline pill
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(5.5), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(30, 41, 59)
    pill.line.color.rgb = c_blue
    tf_pill = pill.text_frame
    p_p = tf_pill.paragraphs[0]
    p_p.text = "GFT × GOOGLE · DESAFIO DE AGENTES DE IA · SMC26"
    p_p.font.size = Pt(11)
    p_p.font.bold = True
    p_p.font.color.rgb = RGBColor(147, 197, 253)

    # Title
    t1_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.8))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "TAMY - Neo Medallion Thesis Monitor"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle
    sub1_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.0), Inches(1.2))
    tf_sub1 = sub1_box.text_frame
    tf_sub1.word_wrap = True
    p_sub1 = tf_sub1.paragraphs[0]
    p_sub1.text = "Monitoramento autônomo de teses de investimento em tempo real com Agentic RAG, Gemini 2.5 e governança Human-in-the-Loop para o Mercado de Capitais."
    p_sub1.font.size = Pt(18)
    p_sub1.font.color.rgb = RGBColor(203, 213, 225)

    # Team Box
    team_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.0), Inches(1.0))
    tf_team = team_box.text_frame
    p_team = tf_team.paragraphs[0]
    p_team.text = "Equipe 05: Neo Medallion  |  Capitão: Edilson Jesus dos Santos Junior (ennt@gft.com)"
    p_team.font.size = Pt(13)
    p_team.font.bold = True
    p_team.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 2: O Problema
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "O Desafio: Sobrecarga Informacional no Mercado de Capitais")

    # Card 1: Volume Explosivo
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = c_card_bg
    c1.line.color.rgb = c_border
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "1. Volume & Ruído\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_accent_red
    p_body = tf_c1.add_paragraph()
    p_body.text = "• Centenas de Fatos Relevantes e notícias diárias na B3.\n• Dificuldade em separar ruído corporativo de eventos de impacto real.\n• Leituras manuais fragmentadas e propensas a sobrecarga cognitiva."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = c_text_dark

    # Card 2: Perda de Premissas
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = c_card_bg
    c2.line.color.rgb = c_border
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "2. Quebras Silenciosas\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_accent_amber
    p_body = tf_c2.add_paragraph()
    p_body.text = "• Teses dependem de premissas complexas (ex: capex, M&A, governança).\n• Mudanças sutis de diretriz invalidam o valuation sem alerta imediato.\n• Dependência exclusiva da memória do analista de research."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = c_text_dark

    # Card 3: Custo de Reação Tardia
    c3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8))
    c3.fill.solid()
    c3.fill.fore_color.rgb = c_card_bg
    c3.line.color.rgb = c_border
    tf_c3 = c3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "3. Risco de Capital\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue
    p_body = tf_c3.add_paragraph()
    p_body.text = "• Reações atrasadas em semanas para realocação de portfólio.\n• Exposição severa a drawdown e perdas patrimoniais.\n• Perda de competitividade frente ao mercado de alta frequência."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = c_text_dark

    # -------------------------------------------------------------
    # SLIDE 3: A Solução TAMY & Arquitetura Multi-Agentes
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "A Solução: Pipeline Multi-Agentes com Agentic RAG")

    steps = [
        ("Agente 1: Triagem", "Gemini 2.5 Flash", "Filtra ruído corporativo e avalia materialidade financeira.", c_blue),
        ("Agente 2: Investigador", "Gemini 2.5 Flash", "Gera busca semântica no Vertex AI Search e Grounding Google.", c_accent_amber),
        ("Agente 3: Analista Líder", "Gemini 2.5 Pro", "Cruza evidências da tese vs notícia, cita trechos e afere risco.", c_accent_red),
        ("Governança & HITL", "Firestore + UI", "Aciona revisão humana em risco ALTO/MUITO ALTO para decisão.", c_accent_green)
    ]

    for idx, (stitle, tech, desc, col) in enumerate(steps):
        left = Inches(0.8 + idx * 2.95)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{stitle}\n"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        
        p_tech = tf.add_paragraph()
        p_tech.text = f"[{tech}]\n"
        p_tech.font.size = Pt(11)
        p_tech.font.bold = True
        p_tech.font.color.rgb = c_text_muted

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = c_text_dark

    # -------------------------------------------------------------
    # SLIDE 4: Diferenciais & Governança Human-in-the-Loop
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "Diferenciais: Grounding Real & Human-in-the-Loop Seletivo")

    # Left Box: HITL
    box_l = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = c_card_bg
    box_l.line.color.rgb = c_border
    tf_l = box_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "🛡️ Human-in-the-Loop Inteligente\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue
    p_body = tf_l.add_paragraph()
    p_body.text = "• Filtro de Ruído Operacional: Alertas com severidade Baixa, Média ou Neutra são arquivados como informativos sem gerar pendência.\n\n• Alertas Críticos (ALTO / MUITO ALTO): Acionam status PENDENTE no feed do analista, exigindo parecer formal e registro de decisão.\n\n• Conformidade e Rastreabilidade: Histórico auditável no Google Cloud Firestore com timestamp, autor e justificativa."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = c_text_dark

    # Right Box: Grounding & Vertex AI Search
    box_r = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = c_card_bg
    box_r.line.color.rgb = c_border
    tf_r = box_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "🌐 Google Search Grounding & Vertex AI Search\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_accent_green
    p_body = tf_r.add_paragraph()
    p_body.text = "• Vertex AI Search: Data Store semântico não-estruturado para indexação profunda de teses de investimento complexas.\n\n• Fallback ChromaDB: Resiliência operacional local em caso de interrupção externa de rede.\n\n• Live Web Grounding: Coleta ativa de notícias da B3 com extração e renderização de URLs reais (ex: Valor Econômico, InfoMoney) para auditoria direta."
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = c_text_dark

    # -------------------------------------------------------------
    # SLIDE 5: Impacto & Resultados de Negócio
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "Impacto de Negócio Mensurável no Mercado de Capitais")

    metrics = [
        ("-95%", "Tempo de Leitura & Triagem", "De horas de varredura manual de Fatos Relevantes para segundos."),
        ("100%", "Rastreabilidade de Decisão", "Cada quebra de premissa fundamentada com citação literal da tese."),
        ("Zero", "Alucinação em Risco Crítico", "Human-in-the-Loop garante validação humana antes de ordens de mesa.")
    ]

    for idx, (m_val, m_title, m_desc) in enumerate(metrics):
        left = Inches(0.8 + idx * 3.95)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = c_light_bg
        card.line.color.rgb = c_blue
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{m_val}\n"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = c_blue
        
        p_t = tf.add_paragraph()
        p_t.text = f"{m_title}\n\n"
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = c_navy

        p_d = tf.add_paragraph()
        p_d.text = m_desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = c_text_dark

    # -------------------------------------------------------------
    # SLIDE 6: Stack & Acesso à Demonstração
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, "Stack Tecnológica & Demonstração Operacional")

    card_stack = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    card_stack.fill.solid()
    card_stack.fill.fore_color.rgb = c_card_bg
    card_stack.line.color.rgb = c_border
    tf_s = card_stack.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "⚙️ Arquitetura Técnica GCP\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_navy
    p_b = tf_s.add_paragraph()
    p_b.text = "• Plataforma de IA: Vertex AI + Google ADK (Agent Development Kit)\n• Modelos: Gemini 2.5 Pro & Gemini 2.5 Flash\n• Vector Search: Vertex AI Search + ChromaDB Fallback\n• Persistência: Google Cloud Firestore + Cloud Storage\n• Backend & Frontend: FastAPI + Streamlit Workstation no Cloud Run"
    p_b.font.size = Pt(13)
    p_b.font.color.rgb = c_text_dark

    card_demo = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    card_demo.fill.solid()
    card_demo.fill.fore_color.rgb = c_card_bg
    card_demo.line.color.rgb = c_border
    tf_d = card_demo.text_frame
    tf_d.word_wrap = True
    p = tf_d.paragraphs[0]
    p.text = "▶️ Como Avaliar o Projeto\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_blue
    p_b2 = tf_d.add_paragraph()
    p_b2.text = "• Demo Cloud Run (Ambiente GCP Corporativo):\n  gcloud run services proxy neomedallion-frontend --port 8501\n\n• Execução Local Autônoma:\n  ./scripts/run_demo.sh (FastAPI + Streamlit)\n\n• Vídeo Pitch + Demo (12m 16s):\n  docs/video/demo_pitch.mov / demo_pitch.mp4\n\n• One-Pager Executivo & Arquitetura (PDF):\n  docs/one-pager.pdf e docs/arquitetura.pdf"
    p_b2.font.size = Pt(13)
    p_b2.font.color.rgb = c_text_dark

    output_path = "docs/apresentacao.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
